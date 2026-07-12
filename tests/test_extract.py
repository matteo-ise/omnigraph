from pathlib import Path

import pytest

from omnigraph.extract import extract_file, get_extractor
from omnigraph.extract.base import ExtractedContent

FIXTURES = Path(__file__).parent / "fixtures"


def test_text_extractor():
    result = extract_file(FIXTURES / "sample.txt")
    assert result is not None
    assert "sample text file" in result.text
    assert result.metadata["type"] == "text"
    assert len(result.chunks) > 0


def test_markdown_extractor():
    result = extract_file(FIXTURES / "sample.md")
    assert result is not None
    assert "markdown" in result.text
    assert result.metadata["type"] == "text"
    assert result.metadata.get("title") == "Test Document"
    assert result.metadata.get("author") == "omnigraph"


def test_code_extractor():
    result = extract_file(FIXTURES / "sample.py")
    assert result is not None
    assert "hello" in result.text
    assert result.metadata["type"] == "code"
    assert result.metadata["language"] == "python"
    assert result.metadata["lines"] > 0


def test_pdf_extractor(tmp_path: Path):
    from pypdf import PdfWriter

    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    pdf_path = tmp_path / "test.pdf"
    with open(pdf_path, "wb") as f:
        writer.write(f)

    result = extract_file(pdf_path)
    assert result is not None
    assert result.metadata["type"] == "pdf"


def test_docx_extractor(tmp_path: Path):
    from docx import Document

    doc = Document()
    doc.add_paragraph("Hello from docx")
    doc.add_paragraph("Second paragraph")
    docx_path = tmp_path / "test.docx"
    doc.save(str(docx_path))

    result = extract_file(docx_path)
    assert result is not None
    assert "Hello from docx" in result.text
    assert result.metadata["type"] == "docx"


def test_xlsx_extractor(tmp_path: Path):
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws["A1"] = "Name"
    ws["B1"] = "Value"
    ws["A2"] = "Test"
    ws["B2"] = 42
    xlsx_path = tmp_path / "test.xlsx"
    wb.save(str(xlsx_path))

    result = extract_file(xlsx_path)
    assert result is not None
    assert "Name" in result.text
    assert "Test" in result.text
    assert result.metadata["type"] == "xlsx"


def test_image_extractor(tmp_path: Path):
    from PIL import Image

    img = Image.new("RGB", (100, 100), color="red")
    img_path = tmp_path / "test.png"
    img.save(str(img_path))

    result = extract_file(img_path)
    assert result is not None
    assert result.metadata["type"] == "image"
    assert result.metadata["width"] == 100
    assert result.metadata["height"] == 100


def test_unknown_extension(tmp_path: Path):
    unknown = tmp_path / "file.xyz123"
    unknown.write_text("some content")
    result = extract_file(unknown)
    assert result is None


def test_get_extractor():
    assert get_extractor(Path("file.txt")) is not None
    assert get_extractor(Path("file.pdf")) is not None
    assert get_extractor(Path("file.py")) is not None
    assert get_extractor(Path("file.xyz123")) is None
