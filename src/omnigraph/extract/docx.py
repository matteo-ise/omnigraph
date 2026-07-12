from __future__ import annotations

from pathlib import Path

from omnigraph.extract.base import ExtractedContent, Extractor


class DocxExtractor(Extractor):
    def extract(self, path: Path) -> ExtractedContent:
        try:
            from docx import Document

            doc = Document(str(path))
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            text = "\n\n".join(paragraphs)

            tables_text = []
            for table in doc.tables:
                for row in table.rows:
                    cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    row_text = " | ".join(cells)
                    if row_text:
                        tables_text.append(row_text)
            if tables_text:
                text += "\n\n" + "\n".join(tables_text)

        except Exception:
            text = ""

        metadata = {
            "type": "docx",
            "extension": ".docx",
            "size": path.stat().st_size,
        }

        chunks = self.chunk_text(text)
        return ExtractedContent(text=text, metadata=metadata, chunks=chunks)
