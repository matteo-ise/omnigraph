from pathlib import Path
from omnigraph.extract.base import ExtractedContent, Extractor

class PptxExtractor(Extractor):
    def extract(self, path: Path) -> ExtractedContent | None:
        try:
            from pptx import Presentation
            prs = Presentation(path)
            text = []
            for slide in prs.slides:
                for shape in slide.shapes:
                    if hasattr(shape, "text"):
                        text.append(shape.text)
            return ExtractedContent(
                text="\n".join(text),
                metadata={"type": "presentation", "title": path.name},
                chunks=[]
            )
        except Exception:
            return None
