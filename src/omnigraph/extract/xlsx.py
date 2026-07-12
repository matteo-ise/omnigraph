from __future__ import annotations

from pathlib import Path

from omnigraph.extract.base import ExtractedContent, Extractor


class XlsxExtractor(Extractor):
    def extract(self, path: Path) -> ExtractedContent:
        try:
            from openpyxl import load_workbook

            wb = load_workbook(str(path), read_only=True, data_only=True)
            sheets_text = []
            for sheet_name in wb.sheetnames:
                ws = wb[sheet_name]
                rows = []
                for row in ws.iter_rows(values_only=True):
                    row_vals = [str(cell) for cell in row if cell is not None]
                    if row_vals:
                        rows.append(" | ".join(row_vals))
                if rows:
                    sheets_text.append(f"[{sheet_name}]\n" + "\n".join(rows))
            wb.close()
            text = "\n\n".join(sheets_text)
        except Exception:
            text = ""

        metadata = {
            "type": "xlsx",
            "extension": ".xlsx",
            "size": path.stat().st_size,
        }

        chunks = self.chunk_text(text)
        return ExtractedContent(text=text, metadata=metadata, chunks=chunks)
